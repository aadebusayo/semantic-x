from __future__ import annotations

import io
from typing import Optional

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def _is_text_blob(name: str) -> bool:
    lower = (name or "").lower()
    return lower.endswith(".txt") or lower.endswith(".md") or lower.endswith(".csv")


def _is_pdf(name: str) -> bool:
    return (name or "").lower().endswith(".pdf")


def _is_docx(name: str) -> bool:
    return (name or "").lower().endswith(".docx")


def _is_xlsx(name: str) -> bool:
    return (name or "").lower().endswith(".xlsx")


def _is_pptx(name: str) -> bool:
    return (name or "").lower().endswith(".pptx")


def _extract_pdf_text(content: bytes) -> Optional[str]:
    reader = PdfReader(io.BytesIO(content))
    texts = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            texts.append(text)
    joined = "\n\n".join(texts).strip()
    return joined or None


def _extract_docx_text(content: bytes) -> Optional[str]:
    document = WordDocument(io.BytesIO(content))
    texts = []

    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            texts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                texts.append(row_text)

    joined = "\n".join(texts).strip()
    return joined or None


def _extract_xlsx_text(content: bytes) -> Optional[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    texts = []

    for sheet in workbook.worksheets:
        sheet_lines = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                sheet_lines.append(" | ".join(values))
        if sheet_lines:
            texts.append(f"[{sheet.title}]\n" + "\n".join(sheet_lines))

    workbook.close()
    joined = "\n\n".join(texts).strip()
    return joined or None


def _extract_pptx_text(content: bytes) -> Optional[str]:
    presentation = Presentation(io.BytesIO(content))
    slides = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            text = text.strip()
            if text:
                parts.append(text)
        if parts:
            slides.append(f"[Slide {index}]\n" + "\n".join(parts))

    joined = "\n\n".join(slides).strip()
    return joined or None


def extract_text(*, blob_name: str, content: bytes) -> Optional[str]:
    if not content:
        return None

    if _is_text_blob(blob_name):
        # Best-effort decode.
        for enc in ("utf-8", "utf-16", "latin-1"):
            try:
                return content.decode(enc, errors="ignore")
            except Exception:
                continue
        return None

    if _is_pdf(blob_name):
        return _extract_pdf_text(content)

    if _is_docx(blob_name):
        return _extract_docx_text(content)

    if _is_xlsx(blob_name):
        return _extract_xlsx_text(content)

    if _is_pptx(blob_name):
        return _extract_pptx_text(content)

    # Unsupported file types are intentionally skipped to keep the service lean.
    return None
