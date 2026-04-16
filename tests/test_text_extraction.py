from __future__ import annotations

import io

from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from services.text_extraction import extract_text


def test_extract_text_reads_docx_content():
    document = WordDocument()
    document.add_heading("Quarterly Review", level=1)
    document.add_paragraph("This document contains the quarterly summary.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Owner"
    table.rows[0].cells[1].text = "Chioma"

    stream = io.BytesIO()
    document.save(stream)

    text = extract_text(blob_name="report.docx", content=stream.getvalue())

    assert text is not None
    assert "Quarterly Review" in text
    assert "This document contains the quarterly summary." in text
    assert "Owner | Chioma" in text


def test_extract_text_reads_xlsx_content():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Month", "Revenue"])
    sheet.append(["April", 42])

    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()

    text = extract_text(blob_name="metrics.xlsx", content=stream.getvalue())

    assert text is not None
    assert "[Metrics]" in text
    assert "Month | Revenue" in text
    assert "April | 42" in text


def test_extract_text_reads_pptx_content():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Launch Review"
    slide.placeholders[1].text = "Traffic increased by 18 percent"

    stream = io.BytesIO()
    presentation.save(stream)

    text = extract_text(blob_name="review.pptx", content=stream.getvalue())

    assert text is not None
    assert "[Slide 1]" in text
    assert "Launch Review" in text
    assert "Traffic increased by 18 percent" in text